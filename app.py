import json, os, re, sys, traceback
from pathlib import Path
from typing import Literal
from pydantic import BaseModel, Field
from PySide6.QtCore import Qt, QThread, Signal
from PySide6.QtGui import QColor
from PySide6.QtWidgets import (QApplication,QMainWindow,QWidget,QVBoxLayout,QHBoxLayout,QLabel,
 QPlainTextEdit,QPushButton,QFileDialog,QListWidget,QListWidgetItem,QTabWidget,QTableWidget,
 QTableWidgetItem,QHeaderView,QMessageBox,QProgressBar,QDialog,QFormLayout,QLineEdit,QComboBox,
 QDialogButtonBox,QSplitter,QAbstractItemView)
from pypdf import PdfReader
from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
import keyring

APP='CV Matcher'; VERSION='0.4'; SERVICE='cv-matcher-ai'

class Requirement(BaseModel):
    id: str
    requirement: str
    priority: Literal['must_have','nice_to_have']
    weight: int = Field(ge=1, le=100)

class VacancyProfile(BaseModel):
    position: str
    requirements: list[Requirement]

class MatchItem(BaseModel):
    requirement_id: str
    status: Literal['MATCH','PARTIAL_MATCH','NO_MATCH','NOT_FOUND']
    evidence: str
    explanation: str

class CandidateAssessment(BaseModel):
    candidate_name: str
    matches: list[MatchItem]
    summary: str

STATUS_FACTOR={'MATCH':1.0,'PARTIAL_MATCH':0.5,'NO_MATCH':0.0,'NOT_FOUND':0.0}
STATUS_ICON={'MATCH':'✓','PARTIAL_MATCH':'~','NO_MATCH':'✗','NOT_FOUND':'?'}

def friendly_api_error(exc):
    text=str(exc)
    low=text.lower()
    if '401' in text or 'authentication' in low or 'invalid api key' in low:
        return 'API-ключ не принят. Проверьте, что ключ скопирован полностью и относится к выбранному провайдеру.'
    if '402' in text or 'insufficient' in low or 'balance' in low or 'quota' in low:
        return 'Недостаточно API-баланса/квоты у провайдера.'
    if '429' in text or 'rate limit' in low:
        return 'Слишком много запросов. Подождите немного и повторите.'
    if 'timeout' in low or 'connection' in low:
        return 'Не удалось подключиться к API. Проверьте интернет, VPN/прокси и повторите.'
    if 'model' in low and ('not found' in low or 'does not exist' in low or 'invalid' in low):
        return 'Эта модель недоступна для данного API-ключа. Нажмите «Обновить список моделей».'
    return text[:700]


def extract_text(path: str) -> str:
    p=Path(path); ext=p.suffix.lower()
    if ext=='.pdf':
        return '\n'.join((page.extract_text() or '') for page in PdfReader(path).pages)
    if ext=='.docx':
        return '\n'.join(x.text for x in Document(path).paragraphs)
    if ext in ('.txt','.md'):
        return p.read_text(encoding='utf-8', errors='ignore')
    raise ValueError(f'Неподдерживаемый формат: {ext}')


DEMO_SKILLS = ['Python','Java','JavaScript','TypeScript','C#','C++','Go','PHP','Ruby','Kotlin','Swift','React','Angular','Vue','Node.js','Django','FastAPI','Flask','Spring','Spring Boot','.NET','SQL','PostgreSQL','MySQL','MongoDB','Redis','Kafka','RabbitMQ','Docker','Kubernetes','AWS','Azure','GCP','Git','Linux','REST','GraphQL','Terraform','Ansible','Jenkins','GitLab CI','CI/CD','Spark','Hadoop','Power BI','Tableau','Excel','1C','SAP','Salesforce','English']

def _contains_term(text, term):
    low=' '+text.lower()+' '
    aliases={'Go':[' golang ',' go '],'C#':['c#','c sharp'],'C++':['c++'],'.NET':['.net','dotnet'],'Node.js':['node.js','nodejs'],'Spring Boot':['spring boot'],'PostgreSQL':['postgresql','postgres'],'Kubernetes':['kubernetes','k8s'],'AWS':['aws','amazon web services'],'GCP':['gcp','google cloud'],'CI/CD':['ci/cd','continuous integration'],'English':['english','английск'],'REST':['rest api',' rest ','restful']}
    return any(x in low for x in aliases.get(term,[term.lower()]))

def demo_profile(vacancy):
    reqs=[]; seen=set(); idx=1
    for skill in DEMO_SKILLS:
        if skill in seen or not _contains_term(vacancy,skill): continue
        seen.add(skill); low=vacancy.lower(); pos=low.find(skill.lower()); context=low[max(0,pos-90):pos+120] if pos>=0 else low
        nice=any(x in context for x in ['желательно','будет плюсом','плюсом','nice to have','preferred','advantage'])
        reqs.append(Requirement(id=f'r{idx}',requirement=skill,priority='nice_to_have' if nice else 'must_have',weight=5 if nice else 10)); idx+=1
    m=re.search(r'(?:от\s*)?(\d+)\s*(?:\+\s*)?(?:лет|года|год|years?|yrs?)',vacancy,re.I)
    if m: reqs.insert(0,Requirement(id='exp',requirement=f'Опыт работы от {m.group(1)} лет',priority='must_have',weight=15))
    if not reqs:
        for ch in [x.strip(' •-\t') for x in re.split(r'[\n;•]+',vacancy) if len(x.strip())>=8][:8]:
            reqs.append(Requirement(id=f'r{idx}',requirement=ch[:120],priority='must_have',weight=10)); idx+=1
    return VacancyProfile(position='Демо-позиция',requirements=reqs[:20])

def demo_assessment(profile, cv_text, filename):
    matches=[]
    for q in profile.requirements:
        if q.id=='exp':
            need=int(re.search(r'(\d+)',q.requirement).group(1)); years=[int(x) for x in re.findall(r'(\d+)\s*(?:\+\s*)?(?:лет|года|год|years?|yrs?)',cv_text,re.I)]
            if years and max(years)>=need: status='MATCH'; ev=f'Указан опыт {max(years)} лет'
            elif years: status='PARTIAL_MATCH'; ev=f'Указан опыт {max(years)} лет'
            else: status='NOT_FOUND'; ev='Срок опыта явно не найден'
        elif _contains_term(cv_text,q.requirement):
            status='MATCH'; ev=next((x.strip() for x in cv_text.splitlines() if _contains_term(x,q.requirement)),q.requirement)[:220]
        else: status='NOT_FOUND'; ev='Упоминание не найдено в резюме'
        matches.append(MatchItem(requirement_id=q.id,status=status,evidence=ev,explanation='Локальная демо-проверка по текстовым совпадениям; AI не использовался.'))
    lines=[x.strip() for x in cv_text.splitlines() if x.strip()]; name=Path(filename).stem
    if lines and len(lines[0])<80 and not any(c.isdigit() for c in lines[0]): name=lines[0]
    ok=sum(m.status=='MATCH' for m in matches)
    return CandidateAssessment(candidate_name=name,matches=matches,summary=f'Демо-режим: найдено явных текстовых совпадений {ok} из {len(matches)}. Это проверка приложения, а не AI-оценка кандидата.')

def calc_score(profile, assessment):
    reqs={r.id:r for r in profile.requirements}; found={m.requirement_id:m for m in assessment.matches}
    total=sum(r.weight for r in profile.requirements) or 1
    earned=sum(r.weight*STATUS_FACTOR.get(found.get(r.id).status if found.get(r.id) else 'NOT_FOUND',0) for r in profile.requirements)
    must=[r for r in profile.requirements if r.priority=='must_have']
    must_ok=sum(1 for r in must if found.get(r.id) and found[r.id].status=='MATCH')
    nice=[r for r in profile.requirements if r.priority=='nice_to_have']
    nice_ok=sum(1 for r in nice if found.get(r.id) and found[r.id].status=='MATCH')
    return round(100*earned/total), must_ok, len(must), nice_ok, len(nice)

class DropList(QListWidget):
    pathsChanged=Signal()
    def __init__(self):
        super().__init__(); self.setAcceptDrops(True); self.setSelectionMode(QAbstractItemView.ExtendedSelection)
        self.setMinimumHeight(150); self.setToolTip('Перетащите сюда PDF, DOCX или TXT')
    def dragEnterEvent(self,e):
        if e.mimeData().hasUrls(): e.acceptProposedAction()
    def dragMoveEvent(self,e): e.acceptProposedAction()
    def dropEvent(self,e):
        self.add_paths([u.toLocalFile() for u in e.mimeData().urls() if u.isLocalFile()]); e.acceptProposedAction()
    def add_paths(self, paths):
        existing={self.item(i).data(Qt.UserRole) for i in range(self.count())}
        for path in paths:
            if Path(path).suffix.lower() in {'.pdf','.docx','.txt','.md'} and path not in existing:
                it=QListWidgetItem(Path(path).name); it.setData(Qt.UserRole,path); it.setToolTip(path); self.addItem(it); existing.add(path)
        self.pathsChanged.emit()
    def paths(self): return [self.item(i).data(Qt.UserRole) for i in range(self.count())]

class ApiTestThread(QThread):
    done=Signal(bool,str)
    def __init__(self, provider, key, model):
        super().__init__(); self.provider=provider; self.key=key; self.model=model
    def run(self):
        try:
            from openai import OpenAI
            base_url='https://api.deepseek.com' if self.provider=='DeepSeek' else None
            client=OpenAI(api_key=self.key, base_url=base_url) if base_url else OpenAI(api_key=self.key)
            if self.provider=='DeepSeek':
                r=client.chat.completions.create(model=self.model,messages=[{'role':'user','content':'Ответь только словом OK.'}],max_tokens=16,extra_body={'thinking':{'type':'disabled'}})
                text=(r.choices[0].message.content or '').strip()
            else:
                r=client.responses.create(model=self.model,input='Ответь только словом OK.',max_output_tokens=32)
                text=(getattr(r,'output_text','') or '').strip()
            self.done.emit(True, f'Подключение успешно. Модель: {self.model}' + (f' · Ответ: {text[:40]}' if text else ''))
        except Exception as e:
            self.done.emit(False, friendly_api_error(e))

class SettingsDialog(QDialog):
    def __init__(self,parent=None):
        super().__init__(parent); self.setWindowTitle('Настройки AI'); self.resize(520,240); f=QFormLayout(self)
        self.provider=QComboBox(); self.provider.addItems(['Демо (без API)','DeepSeek','OpenAI'])
        self.provider.setCurrentText(keyring.get_password(SERVICE,'provider') or 'Демо (без API)')
        self.key=QLineEdit(); self.key.setEchoMode(QLineEdit.Password)
        self.model=QComboBox(); self.model.setEditable(True)
        self.status=QLabel(''); self.status.setWordWrap(True)
        self.test_btn=QPushButton('Проверить подключение'); self.test_btn.clicked.connect(self.test_api)
        self.refresh_btn=QPushButton('Обновить список моделей'); self.refresh_btn.clicked.connect(self.refresh_models)
        self.provider.currentTextChanged.connect(self.provider_changed)
        f.addRow('AI-провайдер:',self.provider); f.addRow('API-ключ:',self.key); f.addRow('Модель:',self.model); f.addRow('',self.refresh_btn); f.addRow('',self.test_btn); f.addRow('',self.status)
        b=QDialogButtonBox(QDialogButtonBox.Save|QDialogButtonBox.Cancel); b.button(QDialogButtonBox.Save).setText('Сохранить'); b.button(QDialogButtonBox.Cancel).setText('Отмена'); b.accepted.connect(self.save); b.rejected.connect(self.reject); f.addRow(b)
        self.provider_changed(self.provider.currentText(), initial=True)
    def provider_changed(self, provider, initial=False):
        self.model.clear()
        if provider=='Демо (без API)':
            self.model.addItem('Локальный demo matcher'); self.key.clear(); self.key.setEnabled(False); self.model.setEnabled(False); self.refresh_btn.setEnabled(False)
            self.status.setText('Бесплатный локальный режим: API-ключ и интернет не нужны. Проверяет интерфейс, Drag & Drop, обработку файлов, таблицу и Excel.'); return
        self.key.setEnabled(True); self.model.setEnabled(True); self.refresh_btn.setEnabled(True)
        if provider=='DeepSeek':
            self.model.addItems(['deepseek-v4-flash','deepseek-v4-pro'])
            key=keyring.get_password(SERVICE,'deepseek_api_key') or ''
            saved=keyring.get_password(SERVICE,'deepseek_model') or 'deepseek-v4-flash'
        else:
            self.model.addItems(['gpt-5.6-luna','gpt-5.6-terra','gpt-5.6-sol'])
            key=keyring.get_password(SERVICE,'openai_api_key') or keyring.get_password('cv-matcher-openai','api_key') or ''
            saved=keyring.get_password(SERVICE,'openai_model') or keyring.get_password('cv-matcher-openai','model') or 'gpt-5.6-luna'
        self.key.setText(key); self.model.setCurrentText(saved); self.status.setText('')
    def refresh_models(self):
        key=self.key.text().strip(); provider=self.provider.currentText()
        if not key:
            QMessageBox.warning(self,'Не хватает данных','Сначала введите API-ключ.'); return
        try:
            from openai import OpenAI
            client=OpenAI(api_key=key, base_url='https://api.deepseek.com') if provider=='DeepSeek' else OpenAI(api_key=key)
            models=sorted({m.id for m in client.models.list().data})
            if provider=='DeepSeek':
                models=[m for m in models if 'deepseek' in m.lower()] or models
            current=self.model.currentText(); self.model.clear(); self.model.addItems(models)
            if current in models: self.model.setCurrentText(current)
            self.status.setText(f'✓ Получено моделей: {len(models)}')
            self.status.setStyleSheet('font-weight:600;color:#188038;')
        except Exception as e:
            self.status.setText('✗ Не удалось получить список моделей: '+friendly_api_error(e))
            self.status.setStyleSheet('font-weight:600;color:#b3261e;')

    def test_api(self):
        key=self.key.text().strip(); model=self.model.currentText().strip(); provider=self.provider.currentText()
        if provider=='Демо (без API)':
            self.status.setText('✓ Демо-режим готов. API не используется.'); self.status.setStyleSheet('font-weight:600;color:#188038;'); return
        if not key or not model: QMessageBox.warning(self,'Не хватает данных','Введите API-ключ и выберите модель.'); return
        self.test_btn.setEnabled(False); self.status.setText('Проверяю подключение…')
        self.tester=ApiTestThread(provider,key,model); self.tester.done.connect(self.test_done); self.tester.start()
    def test_done(self, ok, msg):
        self.test_btn.setEnabled(True); self.status.setText(('✓ ' if ok else '✗ ')+msg)
        self.status.setStyleSheet('font-weight:600;' + ('color:#188038;' if ok else 'color:#b3261e;'))
    def save(self):
        provider=self.provider.currentText(); key=self.key.text().strip(); model=self.model.currentText().strip()
        keyring.set_password(SERVICE,'provider',provider)
        if provider!='Демо (без API)':
            prefix='deepseek' if provider=='DeepSeek' else 'openai'
            keyring.set_password(SERVICE,f'{prefix}_api_key',key); keyring.set_password(SERVICE,f'{prefix}_model',model)
        self.accept()

def llm_structured(client, provider, model, messages, schema_cls):
    if provider=='DeepSeek':
        # DeepSeek's OpenAI-compatible Chat Completions supports JSON output.
        schema=json.dumps(schema_cls.model_json_schema(),ensure_ascii=False)
        msgs=list(messages)
        msgs.insert(0, {'role':'system','content':'Верни только валидный JSON без markdown. JSON должен соответствовать этой JSON Schema: '+schema})
        r=client.chat.completions.create(model=model,messages=msgs,response_format={'type':'json_object'},max_tokens=12000,extra_body={'thinking':{'type':'disabled'}})
        content=r.choices[0].message.content or ''
        if not content.strip(): raise RuntimeError('DeepSeek вернул пустой JSON. Повторите запрос.')
        return schema_cls.model_validate_json(content)
    r=client.responses.parse(model=model,input=messages,text_format=schema_cls)
    return r.output_parsed

class AnalyzeThread(QThread):
    progress=Signal(int,str); done=Signal(object,object); failed=Signal(str)
    def __init__(self,vacancy,paths): super().__init__(); self.vacancy=vacancy; self.paths=paths
    def run(self):
        try:
            from openai import OpenAI
            provider=keyring.get_password(SERVICE,'provider') or 'Демо (без API)'
            if provider=='Демо (без API)':
                self.progress.emit(5,'Демо: извлекаю требования локально…'); profile=demo_profile(self.vacancy); results=[]
                for i,path in enumerate(self.paths):
                    self.progress.emit(10+int(85*i/max(1,len(self.paths))),f'Демо-анализ: {Path(path).name}'); text=extract_text(path)
                    if not text.strip(): raise RuntimeError(f'Не удалось извлечь текст из {Path(path).name}. Возможно, это скан без текстового слоя.')
                    results.append((path,demo_assessment(profile,text,path)))
                self.progress.emit(100,'Демо-анализ готов'); self.done.emit(profile,results); return
            if provider=='DeepSeek':
                key=keyring.get_password(SERVICE,'deepseek_api_key') or os.getenv('DEEPSEEK_API_KEY')
                model=keyring.get_password(SERVICE,'deepseek_model') or 'deepseek-v4-flash'
                if not key: raise RuntimeError('Добавьте DeepSeek API-ключ в Настройки.')
                client=OpenAI(api_key=key, base_url='https://api.deepseek.com')
            else:
                key=keyring.get_password(SERVICE,'openai_api_key') or keyring.get_password('cv-matcher-openai','api_key') or os.getenv('OPENAI_API_KEY')
                model=keyring.get_password(SERVICE,'openai_model') or keyring.get_password('cv-matcher-openai','model') or 'gpt-5.6-luna'
                if not key: raise RuntimeError('Добавьте OpenAI API-ключ в Настройки.')
                client=OpenAI(api_key=key)
            self.progress.emit(3,'Извлекаю требования вакансии…')
            profile=llm_structured(client,provider,model,[
              {'role':'system','content':('Извлеки только профессиональные требования вакансии. Не используй имя, возраст, пол, фото, национальность, семейное положение и другие нерелевантные персональные признаки. '
                 'Каждому требованию дай короткий уникальный id, priority must_have/nice_to_have и вес 1-100 по важности. Не выдумывай требования.')},
              {'role':'user','content':self.vacancy}],VacancyProfile)
            results=[]
            for i,path in enumerate(self.paths):
                self.progress.emit(8+int(85*i/max(1,len(self.paths))),f'Анализ: {Path(path).name}')
                text=extract_text(path)
                if not text.strip(): raise RuntimeError(f'Не удалось извлечь текст из {Path(path).name}. Возможно, это скан без текстового слоя.')
                req_json=json.dumps(profile.model_dump(),ensure_ascii=False)
                prompt=f'''Сравни резюме с требованиями. Для КАЖДОГО requirement_id верни ровно один match.\nСтатусы: MATCH = явно подтверждено; PARTIAL_MATCH = подтверждено частично; NO_MATCH = в CV есть данные, явно противоречащие требованию; NOT_FOUND = данных недостаточно.\nEvidence — короткая цитата/факт только из CV, без выдумок. Не делай выводов по защищенным/нерелевантным персональным признакам. Оценивай только профессиональные данные.\nТРЕБОВАНИЯ:\n{req_json}\n\nРЕЗЮМЕ ({Path(path).name}):\n{text[:90000]}'''
                assessment=llm_structured(client,provider,model,[{'role':'system','content':'Ты аккуратный инструмент сопоставления CV с требованиями. Не принимай решение о найме; только документируй соответствие требованиям.'},{'role':'user','content':prompt}],CandidateAssessment)
                results.append((path,assessment))
            self.progress.emit(100,'Готово'); self.done.emit(profile,results)
        except Exception as e: self.failed.emit(f'{friendly_api_error(e)}\n\nТехнические детали:\n{traceback.format_exc()}')

class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__(); self.setWindowTitle(f'CV Matcher {VERSION}'); self.resize(1150,760); self.profile=None; self.results=[]
        root=QWidget(); self.setCentralWidget(root); lay=QVBoxLayout(root)
        top=QHBoxLayout(); title=QLabel(f'CV MATCHER  <span style="font-size:13px;color:#6b7280">v{VERSION}</span>'); title.setStyleSheet('font-size:24px;font-weight:700'); top.addWidget(title); top.addStretch()
        settings=QPushButton('⚙ Настройки'); settings.clicked.connect(lambda: SettingsDialog(self).exec()); top.addWidget(settings); lay.addLayout(top)
        self.tabs=QTabWidget(); lay.addWidget(self.tabs)
        inp=QWidget(); il=QVBoxLayout(inp); il.addWidget(QLabel('<b>Вакансия</b>'))
        self.vac=QPlainTextEdit(); self.vac.setPlaceholderText('Вставьте сюда описание вакансии…'); il.addWidget(self.vac)
        vh=QHBoxLayout(); vf=QPushButton('Загрузить вакансию из файла'); vf.clicked.connect(self.load_vac); vh.addWidget(vf); demo=QPushButton('★ Загрузить демо-данные'); demo.clicked.connect(self.load_demo); vh.addWidget(demo); vh.addStretch(); il.addLayout(vh)
        il.addWidget(QLabel('<b>Резюме — перетащите файлы в область ниже (Drag & Drop)</b>'))
        self.drop=DropList(); il.addWidget(self.drop)
        bh=QHBoxLayout(); add=QPushButton('+ Добавить резюме'); add.clicked.connect(self.add_cv); rm=QPushButton('Удалить выбранные'); rm.clicked.connect(self.remove_cv); bh.addWidget(add); bh.addWidget(rm); bh.addStretch(); il.addLayout(bh)
        self.go=QPushButton('▶ ПРОАНАЛИЗИРОВАТЬ'); self.go.setMinimumHeight(46); self.go.clicked.connect(self.analyze); il.addWidget(self.go)
        self.progress=QProgressBar(); self.progress.setVisible(False); il.addWidget(self.progress); self.status=QLabel(''); il.addWidget(self.status)
        self.tabs.addTab(inp,'Анализ')
        out=QWidget(); ol=QVBoxLayout(out); self.table=QTableWidget(0,6); self.table.setHorizontalHeaderLabels(['Кандидат','Score','Must have','Nice to have','Файл','Результат']); self.table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch); self.table.doubleClicked.connect(self.show_detail); ol.addWidget(self.table)
        ex=QPushButton('Экспортировать в Excel'); ex.clicked.connect(self.export_excel); ol.addWidget(ex); self.tabs.addTab(out,'Результаты')
    def load_demo(self):
        self.vac.setPlainText("Senior Python Backend Developer\nТребования:\n- Python от 3 лет\n- FastAPI\n- PostgreSQL\n- Docker\n- REST API\n- Git\nБудет плюсом: Kubernetes, Kafka, AWS\nАнглийский B2+")
        demo_dir=Path.home()/'.cv_matcher_demo'; demo_dir.mkdir(exist_ok=True)
        samples={'Anna_Smirnova.txt':"Анна Смирнова\nBackend Developer, 5 лет опыта\nPython, FastAPI, PostgreSQL, Docker, REST API, Git.\nРаботала с Kubernetes и AWS. English B2.",'Ivan_Petrov.txt':"Иван Петров\nPython Developer, 3 года\nPython, Django, PostgreSQL, REST, Git, Docker.\nБазовый опыт Kafka.",'Sergey_Volkov.txt':"Сергей Волков\nSoftware Developer, 2 года\nJava, Spring Boot, MySQL, Git. English B1."}
        paths=[]
        for name,text in samples.items():
            fp=demo_dir/name; fp.write_text(text,encoding='utf-8'); paths.append(str(fp))
        self.drop.clear(); self.drop.add_paths(paths); keyring.set_password(SERVICE,'provider','Демо (без API)')
        QMessageBox.information(self,'Демо готово','Загружены тестовая вакансия и 3 тестовых резюме.\n\nНажмите «ПРОАНАЛИЗИРОВАТЬ». API-ключ не нужен.')

    def load_vac(self):
        p,_=QFileDialog.getOpenFileName(self,'Вакансия','','Documents (*.pdf *.docx *.txt *.md)')
        if p:
            try:self.vac.setPlainText(extract_text(p))
            except Exception as e: QMessageBox.critical(self,'Ошибка',str(e))
    def add_cv(self):
        ps,_=QFileDialog.getOpenFileNames(self,'Резюме','','Documents (*.pdf *.docx *.txt *.md)'); self.drop.add_paths(ps)
    def remove_cv(self):
        for x in self.drop.selectedItems(): self.drop.takeItem(self.drop.row(x))
    def analyze(self):
        if not self.vac.toPlainText().strip() or not self.drop.paths(): QMessageBox.warning(self,'Нужны данные','Добавьте вакансию и хотя бы одно резюме.'); return
        self.go.setEnabled(False); self.progress.setVisible(True); self.progress.setValue(0); self.worker=AnalyzeThread(self.vac.toPlainText(),self.drop.paths()); self.worker.progress.connect(self.on_progress); self.worker.done.connect(self.on_done); self.worker.failed.connect(self.on_fail); self.worker.start()
    def on_progress(self,n,s): self.progress.setValue(n); self.status.setText(s)
    def on_fail(self,s): self.go.setEnabled(True); self.progress.setVisible(False); QMessageBox.critical(self,'Ошибка анализа',s)
    def on_done(self,profile,results):
        self.go.setEnabled(True); self.profile=profile; self.results=results; self.table.setRowCount(0)
        ranked=sorted(results,key=lambda x:calc_score(profile,x[1])[0],reverse=True); self.results=ranked
        for path,a in ranked:
            score,mo,mt,no,nt=calc_score(profile,a); row=self.table.rowCount(); self.table.insertRow(row)
            verdict='Высокое соответствие' if score>=80 else ('Среднее' if score>=60 else 'Низкое')
            vals=[a.candidate_name or Path(path).stem,f'{score}%',f'{mo}/{mt}',f'{no}/{nt}',Path(path).name,verdict]
            for c,v in enumerate(vals): self.table.setItem(row,c,QTableWidgetItem(str(v)))
        self.tabs.setCurrentIndex(1); self.progress.setVisible(False); self.status.setText(f'Обработано резюме: {len(results)}')
    def show_detail(self,index):
        if not self.profile or index.row()>=len(self.results): return
        path,a=self.results[index.row()]; reqs={r.id:r for r in self.profile.requirements}; d=QDialog(self); d.setWindowTitle(a.candidate_name); d.resize(950,650); l=QVBoxLayout(d)
        score,*_=calc_score(self.profile,a); l.addWidget(QLabel(f'<h2>{a.candidate_name} — {score}%</h2><p>{a.summary}</p>'))
        t=QTableWidget(len(a.matches),5); t.setHorizontalHeaderLabels(['Требование','Приоритет','Статус','Evidence','Комментарий']); t.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        for r,m in enumerate(a.matches):
            q=reqs.get(m.requirement_id); vals=[q.requirement if q else m.requirement_id,q.priority if q else '',STATUS_ICON[m.status]+' '+m.status,m.evidence,m.explanation]
            for c,v in enumerate(vals): t.setItem(r,c,QTableWidgetItem(v))
        l.addWidget(t); d.exec()
    def export_excel(self):
        if not self.profile or not self.results: QMessageBox.information(self,'Нет данных','Сначала выполните анализ.'); return
        p,_=QFileDialog.getSaveFileName(self,'Сохранить Excel','cv_match_results.xlsx','Excel (*.xlsx)')
        if not p:return
        wb=Workbook(); ws=wb.active; ws.title='Сводка'; headers=['Кандидат','Score','Must have','Nice to have','Файл']; ws.append(headers)
        for path,a in self.results:
            s,mo,mt,no,nt=calc_score(self.profile,a); ws.append([a.candidate_name,s/100,f'{mo}/{mt}',f'{no}/{nt}',Path(path).name]); ws.cell(ws.max_row,2).number_format='0%'
        for cell in ws[1]: cell.font=Font(bold=True)
        for col in ws.columns: ws.column_dimensions[col[0].column_letter].width=max(14,min(45,max(len(str(x.value or'')) for x in col)+2))
        for path,a in self.results:
            name=re.sub(r'[\\/*?:\[\]]','_',a.candidate_name)[:28] or 'Candidate'; base=name; k=2
            while name in wb.sheetnames: name=f'{base[:25]}_{k}'; k+=1
            sh=wb.create_sheet(name); sh.append(['Требование','Приоритет','Статус','Evidence','Комментарий']); matches={m.requirement_id:m for m in a.matches}
            for q in self.profile.requirements:
                m=matches.get(q.id); sh.append([q.requirement,q.priority,m.status if m else 'NOT_FOUND',m.evidence if m else '',m.explanation if m else ''])
            for cell in sh[1]: cell.font=Font(bold=True)
            sh.freeze_panes='A2'; sh.auto_filter.ref=sh.dimensions
            for col in sh.columns: sh.column_dimensions[col[0].column_letter].width=max(14,min(60,max(len(str(x.value or'')) for x in col)+2))
        wb.save(p); QMessageBox.information(self,'Готово',f'Excel сохранён:\n{p}')

def main():
    app=QApplication(sys.argv); app.setStyle('Fusion'); app.setStyleSheet('QPushButton{padding:6px 10px;} QLineEdit,QComboBox{padding:5px;} QTabWidget::pane{border:1px solid #d1d5db;}'); w=MainWindow(); w.show(); sys.exit(app.exec())
if __name__=='__main__': main()
